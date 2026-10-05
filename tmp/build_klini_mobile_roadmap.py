from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from pathlib import Path


OUT = Path(r"C:\laragon\www\dr-sam-laravel\docs\Plan_de_construccion_Klini_Mobile.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_repeat_heading_rows(table):
    set_repeat_table_header(table.rows[0])


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_keep_with_next(p)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_after = Pt(3)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(text, style="List Number")
    p.paragraph_format.space_after = Pt(4)
    return p


def add_task_table(doc, tasks):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    widths = [Cm(1.7), Cm(5.1), Cm(4.6), Cm(5.7)]
    headers = ["ID", "Tarea", "Entregable", "Criterio de terminado"]
    for idx, (cell, width, header) in enumerate(zip(table.rows[0].cells, widths, headers)):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "17365D")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(8.5)
    set_repeat_heading_rows(table)

    for row_idx, task in enumerate(tasks):
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        for idx, (cell, width, value) in enumerate(zip(cells, widths, task)):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx % 2:
                set_cell_shading(cell, "EEF4FA")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value)
            run.font.size = Pt(8.3)
            if idx == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(23, 54, 93)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


phases = [
    ("Fase 0", "Alcance y diseño funcional", "Backlog aprobado y matriz de funciones por rol"),
    ("Fase 1", "Fundación técnica", "Proyecto Expo reproducible en Android"),
    ("Fase 2", "API móvil Laravel", "API versionada, segura y documentada"),
    ("Fase 3", "Arquitectura de la aplicación", "Navegación, sesión, tema y cliente HTTP"),
    ("Fase 4", "Autenticación", "Ingreso, recuperación y cierre de sesión seguros"),
    ("Fase 5", "Módulo Paciente", "Funciones esenciales del paciente listas"),
    ("Fase 6", "Módulo Médico", "Agenda y atención clínica básica listas"),
    ("Fase 7", "Solicitudes clínicas", "Flujos NPT y oncológico habilitados"),
    ("Fase 8", "Notificaciones y documentos", "Alertas, archivos y tolerancia de red"),
    ("Fase 9", "Calidad y seguridad", "Pruebas, auditoría y aceptación completadas"),
    ("Fase 10", "Beta y publicación", "APK de prueba y proceso de liberación"),
]


task_sections = [
    ("Fase 0 Alcance y diseño funcional", "Definir exactamente qué se construirá antes de programar.", [
        ("KM-001", "Inventariar las funciones web actuales de Paciente y Médico.", "Matriz funcional por rol.", "Cada pantalla y acción tiene decisión: incluir, adaptar o excluir."),
        ("KM-002", "Definir el alcance del producto mínimo viable.", "Lista cerrada del MVP.", "El MVP tiene prioridades y no incluye módulos administrativos."),
        ("KM-003", "Documentar los recorridos principales de ambos perfiles.", "Flujos de usuario.", "Inicio de sesión, navegación y tareas clínicas tienen inicio y fin claros."),
        ("KM-004", "Definir identidad visual de Klini Mobile.", "Colores, tipografía, iconos y componentes base.", "La guía contempla accesibilidad, estados y contraste."),
        ("KM-005", "Definir contratos de datos y reglas de privacidad.", "Catálogo inicial de recursos y datos sensibles.", "Cada dato indica origen, permiso y política de almacenamiento."),
    ]),
    ("Fase 1 Fundación técnica", "Crear una base estable y repetible para el equipo.", [
        ("KM-101", "Crear el repositorio independiente klini-mobile con Expo y TypeScript.", "Proyecto inicial versionado.", "Instala dependencias y abre la pantalla inicial sin errores."),
        ("KM-102", "Configurar Expo Router y rutas tipadas.", "Estructura src/app con grupos Paciente y Médico.", "Las rutas compilan y los grupos permanecen protegidos."),
        ("KM-103", "Configurar ESLint, Prettier y reglas TypeScript estrictas.", "Estándares automáticos de código.", "Lint y comprobación de tipos terminan correctamente."),
        ("KM-104", "Definir ambientes local, pruebas y producción.", "Variables EXPO_PUBLIC_API_URL documentadas.", "No existen URLs ni secretos escritos directamente en el código."),
        ("KM-105", "Configurar Android Studio y un emulador Pixel.", "Emulador operativo.", "La app consume el backend local mediante 10.0.2.2."),
        ("KM-106", "Preparar Expo Go para prototipo y expo-dev-client para el desarrollo real.", "Flujo de ejecución documentado.", "Existe un development build instalable y reproducible."),
    ]),
    ("Fase 2 API móvil Laravel", "Exponer las funciones necesarias sin reutilizar directamente las vistas web.", [
        ("KM-201", "Crear el prefijo versionado /api/mobile/v1.", "Grupo de rutas móvil.", "Todas las rutas móviles están versionadas y aisladas."),
        ("KM-202", "Implementar autenticación móvil con Laravel Sanctum.", "Login y tokens por dispositivo.", "El token tiene capacidades, vencimiento o revocación controlada."),
        ("KM-203", "Crear endpoints /me, logout y revocación de sesiones.", "Gestión de sesión móvil.", "Cerrar sesión invalida el token utilizado."),
        ("KM-204", "Aplicar políticas para roles patient y doctor.", "Autorización por recurso.", "Un paciente no ve datos ajenos y un médico solo accede a pacientes permitidos."),
        ("KM-205", "Definir un formato uniforme de respuestas y errores.", "Contrato JSON común.", "Éxitos, validaciones, permisos y errores usan una estructura consistente."),
        ("KM-206", "Crear endpoints de lectura para Paciente.", "Recursos de perfil, citas, historial, recetas y documentos.", "Cada endpoint está paginado cuando corresponde y tiene pruebas."),
        ("KM-207", "Crear endpoints de lectura y escritura para Médico.", "Agenda, pacientes, notas, recetas y solicitudes.", "Las operaciones respetan permisos, validaciones y auditoría."),
        ("KM-208", "Agregar control de frecuencia, registros y auditoría.", "Protección y trazabilidad.", "Intentos, accesos sensibles y modificaciones quedan registrados."),
        ("KM-209", "Publicar documentación OpenAPI o colección de pruebas.", "Contrato consumible por el equipo móvil.", "Cada endpoint incluye ejemplos, errores y autenticación."),
    ]),
    ("Fase 3 Arquitectura de la aplicación", "Evitar acoplamientos y preparar el crecimiento de la app.", [
        ("KM-301", "Crear el cliente HTTP con URL por ambiente.", "Capa src/services/api.", "Agrega token, tiempo límite y traducción de errores."),
        ("KM-302", "Guardar credenciales con expo-secure-store.", "Almacenamiento seguro del token.", "El token no se guarda en AsyncStorage ni aparece en registros."),
        ("KM-303", "Implementar el estado de autenticación.", "Provider o store de sesión.", "Restaura, renueva y elimina la sesión de forma determinista."),
        ("KM-304", "Crear guardianes de navegación por rol.", "Redirecciones Paciente y Médico.", "No es posible abrir rutas de otro perfil mediante enlaces directos."),
        ("KM-305", "Construir el sistema visual reutilizable.", "Botones, campos, tarjetas, listas y estados.", "Los componentes cubren carga, vacío, error y deshabilitado."),
        ("KM-306", "Configurar caché y consulta de datos.", "Capa de estado remoto.", "Evita solicitudes duplicadas y permite actualización controlada."),
        ("KM-307", "Implementar registro de errores sin datos clínicos.", "Estrategia de observabilidad.", "Los eventos técnicos no contienen tokens ni información médica."),
    ]),
    ("Fase 4 Autenticación", "Entregar el primer recorrido completo y seguro.", [
        ("KM-401", "Diseñar e implementar la pantalla de acceso.", "Login adaptable y accesible.", "Valida campos, muestra errores claros y evita envíos duplicados."),
        ("KM-402", "Implementar carga inicial y restauración de sesión.", "Pantalla de arranque.", "La app decide correctamente entre login y panel del usuario."),
        ("KM-403", "Implementar recuperación y cambio de contraseña.", "Flujo de recuperación.", "No revela si una cuenta existe y respeta políticas del backend."),
        ("KM-404", "Implementar cierre de sesión y expiración.", "Salida segura.", "Revoca el token, limpia datos locales y regresa al login."),
        ("KM-405", "Probar accesos cruzados y sesiones inválidas.", "Pruebas de seguridad de sesión.", "Los accesos indebidos reciben 401 o 403 y no filtran información."),
    ]),
    ("Fase 5 Módulo Paciente", "Trasladar las funciones esenciales del portal del paciente.", [
        ("KM-501", "Crear el panel principal del paciente.", "Resumen de próximas citas, recetas y avisos.", "Los datos coinciden con Laravel y permiten navegación directa."),
        ("KM-502", "Crear perfil y datos personales.", "Consulta y edición permitida del perfil.", "Los campos editables y de solo lectura respetan las reglas actuales."),
        ("KM-503", "Crear agenda y detalle de citas.", "Listado, filtros y detalle.", "Estados, fechas, médico y unidad se muestran correctamente."),
        ("KM-504", "Crear historial clínico.", "Vista cronológica y detalle.", "Solo se muestran registros autorizados y los vacíos son comprensibles."),
        ("KM-505", "Crear recetas y medicamentos.", "Listado y detalle de prescripciones.", "Incluye vigencia, indicaciones y estado sin permitir acciones indebidas."),
        ("KM-506", "Crear estudios y documentos.", "Listado y visualización segura.", "Las descargas requieren autenticación y manejan archivos no disponibles."),
        ("KM-507", "Crear notificaciones del paciente.", "Bandeja y marcado como leído.", "Contadores y estados se sincronizan con el backend."),
    ]),
    ("Fase 6 Módulo Médico", "Permitir que el médico atienda su agenda y pacientes autorizados.", [
        ("KM-601", "Crear el panel principal del médico.", "Resumen de agenda, pendientes y avisos.", "Los indicadores corresponden exclusivamente al médico autenticado."),
        ("KM-602", "Crear agenda diaria y calendario.", "Agenda navegable.", "Permite filtrar, abrir citas y actualizar solo acciones autorizadas."),
        ("KM-603", "Crear catálogo de pacientes autorizados.", "Buscador y listado.", "No devuelve pacientes sin relación clínica vigente."),
        ("KM-604", "Crear expediente resumido del paciente.", "Datos clínicos por secciones.", "Cada acceso valida relación, permiso y queda auditado."),
        ("KM-605", "Crear notas y registro de consulta.", "Formulario clínico validado.", "Guarda borrador o registro final según las reglas del backend."),
        ("KM-606", "Crear prescripciones médicas.", "Alta y consulta de recetas.", "Medicamentos, dosis e indicaciones usan catálogos y validaciones."),
        ("KM-607", "Crear seguimiento de solicitudes.", "Listado de solicitudes y estados.", "Los estados coinciden con Dr. Sam y se actualizan al refrescar."),
    ]),
    ("Fase 7 Solicitudes clínicas", "Llevar a móvil los flujos clínicos complejos después de estabilizar ambos módulos.", [
        ("KM-701", "Adaptar el formulario nutricional NPT.", "Formulario móvil por pasos.", "Valida datos clínicos, componentes, cantidades y fecha de entrega."),
        ("KM-702", "Adaptar el formulario oncológico.", "Formulario móvil por pasos.", "Soporta mezclas, medicamentos, dosis, diluyentes, vías y fechas."),
        ("KM-703", "Integrar catálogos y prevalidación.", "Selección y validación remota.", "No permite enviar productos incompatibles o no disponibles."),
        ("KM-704", "Implementar adjuntos clínicos.", "Carga segura de documentos.", "Valida formato, tamaño, progreso y errores de transferencia."),
        ("KM-705", "Implementar confirmación y envío idempotente.", "Resumen previo al envío.", "Un doble toque no genera solicitudes duplicadas."),
        ("KM-706", "Mostrar seguimiento y observaciones.", "Detalle de estado.", "Presenta autorizaciones, integración, observaciones y resultado operativo."),
    ]),
    ("Fase 8 Notificaciones documentos y red", "Preparar la aplicación para condiciones reales de uso.", [
        ("KM-801", "Configurar notificaciones push en development build.", "Registro de dispositivo y recepción de alertas.", "El usuario puede habilitar o revocar notificaciones."),
        ("KM-802", "Implementar deep links desde notificaciones.", "Navegación a la pantalla relacionada.", "El enlace valida sesión, rol y permisos antes de abrir datos."),
        ("KM-803", "Implementar visualización y descarga de archivos.", "Visor y manejo de descargas.", "Los archivos temporales se eliminan y no quedan expuestos."),
        ("KM-804", "Definir comportamiento sin conexión.", "Mensajes, caché permitida y reintentos.", "La app distingue entre datos guardados, datos vencidos y acciones no enviadas."),
        ("KM-805", "Manejar cambios de red y reanudación.", "Recuperación automática controlada.", "No duplica acciones al reconectarse ni pierde formularios en curso."),
    ]),
    ("Fase 9 Calidad y seguridad", "Demostrar que la aplicación funciona y protege la información clínica.", [
        ("KM-901", "Crear pruebas unitarias de servicios, stores y validadores.", "Suite unitaria móvil.", "Cubre autenticación, errores, formatos y reglas críticas."),
        ("KM-902", "Crear pruebas de componentes y navegación.", "Suite de interfaz.", "Cubre estados de carga, vacío, error y permisos."),
        ("KM-903", "Crear pruebas de integración de la API móvil.", "Suite Laravel móvil.", "Verifica 401, 403, validaciones, propiedad y auditoría."),
        ("KM-904", "Ejecutar pruebas extremo a extremo en Android.", "Recorridos automatizados o guiones reproducibles.", "Paciente y Médico completan sus recorridos principales."),
        ("KM-905", "Realizar revisión de seguridad y privacidad.", "Lista de hallazgos resueltos.", "No hay secretos, registros clínicos indebidos ni acceso horizontal."),
        ("KM-906", "Realizar pruebas de accesibilidad y dispositivos.", "Matriz de compatibilidad.", "Textos, contraste, teclado, escalado y tamaños de pantalla son utilizables."),
        ("KM-907", "Ejecutar aceptación con usuarios.", "Acta de aceptación y pendientes.", "Paciente y Médico validan tareas críticas con datos de prueba."),
    ]),
    ("Fase 10 Beta y publicación", "Entregar una versión instalable, observable y recuperable.", [
        ("KM-1001", "Configurar nombre, identificadores, icono y splash de Klini Mobile.", "Configuración nativa final.", "Se visualiza correctamente en el development build y release."),
        ("KM-1002", "Configurar perfiles EAS para desarrollo, vista previa y producción.", "eas.json y credenciales controladas.", "Cada perfil genera el artefacto esperado sin exponer secretos."),
        ("KM-1003", "Generar APK de distribución interna.", "Versión beta instalable.", "Se instala en un dispositivo limpio y consume el ambiente de pruebas."),
        ("KM-1004", "Preparar monitoreo y procedimiento de soporte.", "Guía operativa.", "Incluye diagnóstico, revocación de sesiones y respuesta ante incidentes."),
        ("KM-1005", "Preparar ficha y política para Google Play.", "Materiales de publicación.", "Permisos, privacidad, capturas y clasificación están completos."),
        ("KM-1006", "Publicar beta cerrada y recopilar resultados.", "Versión candidata.", "No existen defectos críticos y los pendientes están priorizados."),
    ]),
]


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.12

for style_name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 13)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(size)
    style.font.bold = True
    if style_name == "Title":
        title_style_ppr = style.element.get_or_add_pPr()
        title_border = title_style_ppr.find(qn("w:pBdr"))
        if title_border is not None:
            title_style_ppr.remove(title_border)
styles["Heading 1"].paragraph_format.space_before = Pt(14)
styles["Heading 1"].paragraph_format.space_after = Pt(7)
styles["Heading 2"].paragraph_format.space_before = Pt(10)
styles["Heading 2"].paragraph_format.space_after = Pt(5)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("Plan de construcción de Klini Mobile")
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
r = subtitle.add_run("Hoja de ruta por tareas para los módulos Paciente y Médico")
r.bold = True
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(50, 50, 50)

intro = doc.add_paragraph(
    "Este documento organiza la construcción de Klini Mobile en tareas verificables. "
    "La aplicación será un proyecto Expo y React Native independiente, mientras que Dr. Sam continuará como backend y fuente oficial de datos. "
    "La secuencia prioriza seguridad, trazabilidad y entregas pequeñas que puedan probarse antes de avanzar."
)
intro.paragraph_format.space_after = Pt(10)

p = doc.add_paragraph()
r = p.add_run("Alcance inicial. ")
r.bold = True
p.add_run("Únicamente los módulos Paciente y Médico, con sus funciones autorizadas. Se excluyen administración, operación hospitalaria, inventarios, instituciones, seguros y configuración general.")

add_heading(doc, "Cómo usar esta hoja de ruta", 1)
for item in [
    "Trabajar las fases en orden y no iniciar una tarea cuyo requisito previo esté pendiente.",
    "Marcar una tarea como terminada solamente cuando cumpla su criterio de aceptación y tenga evidencia de prueba.",
    "Registrar decisiones técnicas y cambios de alcance junto al identificador de la tarea.",
    "Mantener datos clínicos reales fuera de ambientes de desarrollo y demostración.",
]:
    add_bullet(doc, item)

add_heading(doc, "Tecnología acordada", 1)
tech = doc.add_table(rows=1, cols=2)
tech.style = "Table Grid"
tech.alignment = WD_TABLE_ALIGNMENT.CENTER
tech.autofit = False
for idx, (cell, text, width) in enumerate(zip(tech.rows[0].cells, ["Componente", "Decisión"], [Cm(5), Cm(12.1)])):
    cell.width = width
    set_cell_shading(cell, "17365D")
    set_cell_margins(cell)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = cell.paragraphs[0].add_run(text)
    rr.bold = True
    rr.font.color.rgb = RGBColor(255, 255, 255)
for idx, row in enumerate([
    ("Aplicación", "Expo, React Native y TypeScript"),
    ("Navegación", "Expo Router con rutas separadas por rol"),
    ("Backend", "Laravel API versionada dentro de Dr. Sam"),
    ("Autenticación", "Laravel Sanctum y token por dispositivo"),
    ("Seguridad local", "expo-secure-store para credenciales"),
    ("Desarrollo", "Expo Go para prototipo y expo-dev-client para el proyecto real"),
    ("Android", "Android Studio y emulador Pixel"),
    ("Iconos", "@expo/vector-icons"),
]):
    table_row = tech.add_row()
    set_row_cant_split(table_row)
    cells = table_row.cells
    if idx % 2:
        for c in cells:
            set_cell_shading(c, "EEF4FA")
    for c, value, width in zip(cells, row, [Cm(5), Cm(12.1)]):
        c.width = width
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(c)
        c.paragraphs[0].add_run(value).font.size = Pt(9)

doc.add_page_break()
add_heading(doc, "Mapa general de fases", 1)
summary = doc.add_table(rows=1, cols=3)
summary.style = "Table Grid"
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
summary.autofit = False
sum_widths = [Cm(2.2), Cm(6.2), Cm(8.7)]
for cell, width, header in zip(summary.rows[0].cells, sum_widths, ["Fase", "Enfoque", "Resultado"]):
    cell.width = width
    set_cell_shading(cell, "17365D")
    set_cell_margins(cell)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = cell.paragraphs[0].add_run(header)
    rr.bold = True
    rr.font.color.rgb = RGBColor(255, 255, 255)
    rr.font.size = Pt(9)
set_repeat_heading_rows(summary)
for idx, row in enumerate(phases):
    table_row = summary.add_row()
    set_row_cant_split(table_row)
    cells = table_row.cells
    if idx % 2:
        for c in cells:
            set_cell_shading(c, "EEF4FA")
    for j, (c, value, width) in enumerate(zip(cells, row, sum_widths)):
        c.width = width
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(c)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
        rr = c.paragraphs[0].add_run(value)
        rr.font.size = Pt(8.7)
        if j == 0:
            rr.bold = True

doc.add_page_break()
add_heading(doc, "Backlog de construcción", 1)
p = doc.add_paragraph("Cada fila corresponde a una tarea rastreable. Los identificadores deben utilizarse en commits, incidencias y reportes de avance.")
p.paragraph_format.space_after = Pt(10)

for index, (heading, objective, tasks) in enumerate(task_sections):
    add_heading(doc, heading, 2)
    p = doc.add_paragraph()
    r = p.add_run("Objetivo. ")
    r.bold = True
    p.add_run(objective)
    p.paragraph_format.space_after = Pt(6)
    add_task_table(doc, tasks)

add_heading(doc, "Hitos de entrega", 1)
milestones = [
    ("Hito 1 Base instalada", "Fases 0 y 1", "La app abre en Android y el equipo comparte una configuración reproducible."),
    ("Hito 2 Sesión segura", "Fases 2 a 4", "Paciente y Médico ingresan, recuperan sesión y solo ven rutas permitidas."),
    ("Hito 3 MVP Paciente", "Fase 5", "El paciente consulta su información, citas, historial, recetas y documentos."),
    ("Hito 4 MVP Médico", "Fase 6", "El médico consulta agenda y pacientes, registra atención y genera recetas."),
    ("Hito 5 Clínica ampliada", "Fases 7 y 8", "Solicitudes clínicas, documentos y notificaciones funcionan en condiciones reales."),
    ("Hito 6 Beta cerrada", "Fases 9 y 10", "Existe una versión Android aprobada para usuarios de prueba."),
]
milestone_table = doc.add_table(rows=1, cols=3)
milestone_table.style = "Table Grid"
milestone_table.alignment = WD_TABLE_ALIGNMENT.CENTER
milestone_table.autofit = False
milestone_widths = [Cm(4), Cm(3), Cm(10.1)]
for cell, width, header in zip(milestone_table.rows[0].cells, milestone_widths, ["Hito", "Cobertura", "Condición de salida"]):
    cell.width = width
    set_cell_shading(cell, "17365D")
    set_cell_margins(cell)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = cell.paragraphs[0].add_run(header)
    rr.bold = True
    rr.font.color.rgb = RGBColor(255, 255, 255)
    rr.font.size = Pt(9)
for idx, row in enumerate(milestones):
    table_row = milestone_table.add_row()
    set_row_cant_split(table_row)
    cells = table_row.cells
    if idx % 2:
        for c in cells:
            set_cell_shading(c, "EEF4FA")
    for j, (c, value, width) in enumerate(zip(cells, row, milestone_widths)):
        c.width = width
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(c)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 1 else WD_ALIGN_PARAGRAPH.LEFT
        c.paragraphs[0].add_run(value).font.size = Pt(8.7)

add_heading(doc, "Primera iteración recomendada", 1)
doc.add_paragraph("Para comenzar sin dispersar el esfuerzo, la primera iteración debe terminar con una aplicación instalada que autentique ambos perfiles contra Laravel.")
for step in [
    "Completar KM-001 a KM-005 y aprobar el alcance del MVP.",
    "Completar KM-101 a KM-106 y ejecutar Klini Mobile en un emulador Pixel.",
    "Completar KM-201 a KM-205 para contar con autenticación y permisos básicos.",
    "Completar KM-301 a KM-305 para tener navegación, sesión y componentes base.",
    "Completar KM-401, KM-402, KM-404 y KM-405 para validar el recorrido de acceso.",
]:
    add_number(doc, step)

add_heading(doc, "Definición general de terminado", 1)
for item in [
    "El código compila, pasa lint, comprobación de tipos y pruebas relacionadas.",
    "La función se probó en un emulador Android y, cuando aplique, en un dispositivo físico.",
    "Los estados de carga, vacío, error, permiso denegado y sesión vencida están resueltos.",
    "La API valida autenticación, autorización, propiedad de los datos y entradas del usuario.",
    "No se registran tokens, contraseñas ni información clínica sensible en consola o telemetría.",
    "La tarea cuenta con evidencia de prueba y documentación mínima para reproducirla.",
]:
    add_bullet(doc, item)

add_heading(doc, "Referencias técnicas", 1)
refs = [
    "Expo Router: https://docs.expo.dev/router/introduction/",
    "Development builds: https://docs.expo.dev/develop/development-builds/introduction/",
    "Expo SecureStore: https://docs.expo.dev/versions/latest/sdk/securestore/",
    "Laravel Sanctum: https://laravel.com/docs/sanctum",
]
for ref in refs:
    add_bullet(doc, ref)

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.text = "Klini Mobile  Plan de construcción"
footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
footer_p.runs[0].font.size = Pt(8)
footer_p.runs[0].font.color.rgb = RGBColor(100, 100, 100)
page_p = footer.add_paragraph()
add_page_number(page_p)

doc.core_properties.title = "Plan de construcción de Klini Mobile"
doc.core_properties.subject = "Hoja de ruta por tareas para los módulos Paciente y Médico"
doc.core_properties.author = "Equipo Klini"
doc.core_properties.keywords = "Klini Mobile, Expo, React Native, Laravel, Paciente, Médico"
doc.save(OUT)
print(OUT)

